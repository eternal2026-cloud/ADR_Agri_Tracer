"""
verificar_estudio_top.py — Verificación independiente del Excel del Estudio de cosechador TOP.

Lee SOLO las hojas de datos del Excel que descarga Integra (Evaluaciones, Tomas, Estándares
(detalle), REBA, Tablas REBA, Ergonomía (respuestas), Encuesta), rehace cada KPI en Python sin
usar ninguna fórmula del libro y lo compara con la columna «Valor en Integra» de la hoja KPIs.
También rehace el tiempo estándar de cada actividad y el puntaje REBA de cada postura.

Uso:
    pip install openpyxl
    python verificar_estudio_top.py Estudio_TOP_2026-10-29.xlsx

Sale con código 0 si todo coincide y 1 si algo no coincide.
"""
import statistics
import sys

from openpyxl import load_workbook


def filas(wb, hoja):
    """Filas de una hoja como diccionarios {encabezado: valor}."""
    if hoja not in wb.sheetnames:
        return []
    it = wb[hoja].iter_rows(values_only=True)
    cab = next(it, None) or []
    return [dict(zip(cab, f)) for f in it if any(v not in (None, '') for v in f)]


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def percentil(orden, q):
    """Igual que TOP.cuantil de la app y PERCENTILE de Excel (interpolación lineal)."""
    pos = (len(orden) - 1) * q
    b = int(pos)
    r = pos - b
    return orden[b] + r * (orden[b + 1] - orden[b]) if b + 1 < len(orden) else orden[b]


def main(ruta):
    wb = load_workbook(ruta, data_only=False)
    par = {f['Parámetro']: f['Valor'] for f in filas(wb, 'Parámetros')}
    sup = float(par.get('Suplemento % (fatiga y necesidades)') or 0)

    evals = filas(wb, 'Evaluaciones')
    tomas = filas(wb, 'Tomas')
    cump = filas(wb, 'Estándares (detalle)')
    reba = filas(wb, 'REBA')
    ergo = filas(wb, 'Ergonomía (respuestas)')
    enc = filas(wb, 'Encuesta')

    validas = [t for t in tomas if num(t['Duración (s)']) is not None and t['Duración (s)'] >= 0]
    clases = ['Productiva', 'Soporte necesario', 'Espera / eliminable', 'Personal']
    por_eval = {}
    for t in validas:
        d = por_eval.setdefault(t['Código'], dict.fromkeys(clases, 0))
        d[t['Clase']] = d.get(t['Clase'], 0) + t['Duración (s)']
    total = sum(sum(d.values()) for d in por_eval.values())
    pct = lambda clase: sum(d[clase] for d in por_eval.values()) / total * 100 if total else 0

    # Tiempo estándar por actividad: mediana × (1 + suplemento), redondeado al segundo.
    problemas = []
    estandares = {}
    for t in validas:
        estandares.setdefault(t['ID actividad'], []).append(t['Duración (s)'])
    integra_te = {f['ID']: f for f in filas(wb, 'Tiempo estándar')}
    for act, v in sorted(estandares.items()):
        v.sort()
        med = statistics.median(v)
        est = round(med * (1 + sup / 100) + 1e-9)
        fila = integra_te.get(act)
        if fila is None or abs((num(fila['Integra: estándar (s)']) or -1) - est) > 0.5:
            problemas.append(f'Tiempo estándar actividad {act}: Python {est} s vs Integra {fila and fila["Integra: estándar (s)"]}')
        p25, p75 = percentil(v, .25), percentil(v, .75)
        print(f'  actividad {act:>3}: {len(v):>3} tomas · mediana {med:7.1f} s · P25 {p25:7.1f} · P75 {p75:7.1f} · estándar {est} s')

    # REBA: se rehace cada postura con las tablas A, B y C de la hoja «Tablas REBA».
    tr_hoja = wb['Tablas REBA'] if 'Tablas REBA' in wb.sheetnames else None
    malas_reba = 0
    if tr_hoja is not None:
        ta = [[tr_hoja.cell(3 + i, 2 + j).value for j in range(4)] for i in range(15)]
        tb = [[tr_hoja.cell(3 + i, 8 + j).value for j in range(3)] for i in range(12)]
        tc = [[tr_hoja.cell(3 + i, 13 + j).value for j in range(12)] for i in range(12)]
        for r in reba:
            g = lambda k: num(r.get(k)) or 0
            try:
                cu, tro = min(3, g('Cuello') + g('Cuello +1 (1=sí)')), min(5, g('Tronco') + g('Tronco +1'))
                pi = min(4, g('Piernas') + g('Piernas +'))
                br = max(1, min(6, g('Brazo') + g('Hombro +1') - g('Apoyo −1')))
                mu = min(3, g('Muñeca') + g('Muñeca +1'))
                a = ta[(tro - 1) * 3 + cu - 1][pi - 1]
                b = tb[(br - 1) * 2 + g('Antebrazo') - 1][mu - 1]
                pa, pb = min(12, a + g('Carga') + g('Brusca +1')), min(12, b + g('Agarre'))
                final = tc[pa - 1][pb - 1] + g('Estática +1') + g('Repetitiva +1') + g('Inestable +1')
            except (IndexError, TypeError):
                final = None
            if final != r['REBA final']:
                malas_reba += 1
                problemas.append(f'REBA {r["Código"]} ({r["Actividad"]}): Python {final} vs guardado {r["REBA final"]}')

    c = sum(1 for x in cump if x['Resultado'] == 'Cumple')
    nc = sum(1 for x in cump if x['Resultado'] == 'No cumple')
    finales = [x['REBA final'] for x in reba if num(x['REBA final']) is not None]
    calc = {
        'Evaluaciones en el filtro': len(evals),
        'Cosechadores TOP distintos': len({e['Cosechador'] for e in evals if e['Cosechador']}),
        'Tomas de tiempo registradas': len(tomas),
        'Tomas válidas (con fin)': len(validas),
        'Jornadas medidas': sum(1 for d in por_eval.values() if sum(d.values()) > 0),
        '% del tiempo productivo': pct('Productiva'),
        '% del tiempo en espera / eliminable': pct('Espera / eliminable'),
        'Actividades con tiempo estándar': len(estandares),
        'Actividades cuyo estándar coincide con Integra': len(estandares) - sum(1 for p in problemas if p.startswith('Tiempo estándar')),
        'Cumplimiento global de estándares %': c / (c + nc) * 100 if c + nc else None,
        'Posturas REBA evaluadas': len(reba),
        'REBA promedio': statistics.mean(finales) if finales else None,
        'REBA máximo': max(finales) if finales else None,
        '% posturas en riesgo alto o muy alto': sum(1 for x in reba if x['Riesgo'] in ('Alto', 'Muy alto')) / len(reba) * 100 if reba else None,
        'Posturas REBA recalculadas que coinciden': len(reba) - malas_reba,
        'Entrevistas ergonómicas': len({x['Código'] for x in ergo}),
        'Encuestas aceptadas': sum(1 for x in enc if x['Aceptó'] == 'Sí'),
    }

    print('\nKPI · Python · Integra')
    for f in filas(wb, 'KPIs'):
        k, integra = f.get('KPI'), f.get('Valor en Integra')
        if k not in calc:
            continue
        py = calc[k]
        ok = (py is None and integra in (None, '')) or (py is not None and num(integra) is not None and abs(py - integra) <= 0.1)
        print(f'  {"OK " if ok else "MAL"} {k}: {py if py is None else round(py, 2)} · {integra if integra is None else round(integra, 2)}')
        if not ok:
            problemas.append(f'KPI «{k}»: Python {py} vs Integra {integra}')

    print('\n' + ('Todo coincide.' if not problemas else f'{len(problemas)} diferencia(s):'))
    for p in problemas:
        print('  - ' + p)
    return 0 if not problemas else 1


if __name__ == '__main__':
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1]))
